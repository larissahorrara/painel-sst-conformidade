"""Exporta cada tabela de data/base-dados-sst.xlsx para data/csv/<tabela>.csv.

Os CSVs usam vírgula como separador, ponto como separador decimal e datas no formato AAAA-MM-DD,
o padrão aceito pelo Power BI, pelo Python (pandas) e pela pré-visualização de tabelas do GitHub.

O arquivo Excel precisa ter os valores das fórmulas calculados: abra e salve no Excel
(ou no LibreOffice) antes de exportar, se ele acabou de ser gerado por scripts/montar_base.py.

Uso (na raiz do repositório):  python scripts/exportar_csv.py
"""
import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / 'data' / 'base-dados-sst.xlsx'
DESTINO = RAIZ / 'data' / 'csv'


def formatar(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float):
        return f'{valor:.6f}'.rstrip('0').rstrip('.')
    return '' if valor is None else valor


wb = load_workbook(ORIGEM, data_only=True)
DESTINO.mkdir(parents=True, exist_ok=True)
total = 0
for ws in wb.worksheets:
    for tabela in ws.tables.values():
        linhas = [[formatar(c.value) for c in linha] for linha in ws[tabela.ref]]
        with open(DESTINO / f'{tabela.displayName}.csv', 'w', newline='', encoding='utf-8') as f:
            csv.writer(f).writerows(linhas)
        total += 1
print(f'{total} tabelas exportadas para {DESTINO.relative_to(RAIZ)}/')
