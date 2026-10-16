"""Monta a base de dados (Excel) do Painel de SST e Conformidade a partir de data/dados.json.

Uso (na raiz do repositório):
    pip install -r requirements.txt
    python scripts/montar_base.py data/dados.json data/base-dados-sst.xlsx

As colunas de fórmula são recalculadas automaticamente quando o arquivo é aberto no Excel.
"""
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ENTRADA, SAIDA = sys.argv[1], sys.argv[2]
J = json.load(open(ENTRADA))
REF = date.fromisoformat(J['hoje'])
REF_CELULA = "'LEIA-ME'!$B$4"
UNIDADES = J['unidades']
PU = J['porUnidade']
SIGLA = {u['id']: u['sigla'] for u in UNIDADES}

FONTE = Font(name='Arial', size=10)
FONTE_NEGRITO = Font(name='Arial', size=10, bold=True)
FONTE_TITULO = Font(name='Arial', size=14, bold=True, color='1E3A5F')
FONTE_ENTRADA = Font(name='Arial', size=10, color='0000FF')  # valores digitados (convenção de planilhas)
FMT_DATA, FMT_PCT, FMT_INT, FMT_DEC1, FMT_DEC2 = 'dd/mm/yyyy', '0.0%', '#,##0', '#,##0.0', '#,##0.00'

d = lambda s: date.fromisoformat(s)
wb = Workbook()
DICIONARIO = []  # (tabela, coluna, tipo, descrição, relacionamento)


def tabela(nome, descricao, colunas, linhas, larguras=None):
    """Cria uma aba com uma Tabela do Excel. colunas: [(nome, tipo, descrição, formato, relacionamento)].
    Em `linhas`, um valor str que começa com '=' é fórmula; '{r}' é trocado pelo número da linha."""
    ws = wb.create_sheet(nome)
    for c, (col, *_rest) in enumerate(colunas, 1):
        ws.cell(row=1, column=c, value=col).font = FONTE_NEGRITO
    for i, linha in enumerate(linhas, 2):
        for c, valor in enumerate(linha, 1):
            if isinstance(valor, str) and valor.startswith('='):
                valor = valor.replace('{r}', str(i))
            cel = ws.cell(row=i, column=c, value=valor)
            cel.font = FONTE
            fmt = colunas[c - 1][3]
            if fmt:
                cel.number_format = fmt
    ultima = get_column_letter(len(colunas))
    t = Table(displayName=nome, ref=f'A1:{ultima}{max(2, len(linhas) + 1)}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = 'A2'
    for c, (col, *_rest) in enumerate(colunas, 1):
        largura = (larguras or {}).get(col)
        if not largura:
            amostra = [len(str(col))] + [len(str(l[c - 1])) for l in linhas[:200] if not str(l[c - 1]).startswith('=')]
            largura = min(48, max(10, max(amostra) + 2))
        ws.column_dimensions[get_column_letter(c)].width = largura
    for col, tipo, desc, _fmt, rel in colunas:
        DICIONARIO.append((nome, col, tipo, desc, rel or ''))
    TABELAS.append((nome, descricao, len(linhas)))
    return ws


TABELAS = []
UN = ('id_unidade', 'Texto', 'Sigla da unidade', None, 'dUnidade[id_unidade]')

# =====================================================================
# DIMENSÕES
# =====================================================================
linhas = []
for u in UNIDADES:
    tipo, resto = u['nome'].split(' ', 1)          # "Matriz – Brasília/DF" ou "Filial São Paulo/SP"
    cidade, uf = resto.lstrip('– ').split('/')
    linhas.append([u['sigla'], u['nome'], tipo, cidade, uf, u['colaboradores'], '=F{r}*%d' % J['horasAno'], J['sesmtConfig'][u['id']]['grau']])
tabela('dUnidade', 'Unidades (matriz e filiais), efetivo e horas-homem trabalhadas', [
    ('id_unidade', 'Texto (chave)', 'Sigla da unidade', None, ''),
    ('nome_unidade', 'Texto', 'Nome completo da unidade', None, ''),
    ('tipo', 'Texto', 'Matriz ou Filial', None, ''),
    ('cidade', 'Texto', 'Cidade', None, ''),
    ('uf', 'Texto', 'Unidade da Federação', None, ''),
    ('colaboradores', 'Inteiro', 'Efetivo atual', FMT_INT, ''),
    ('horas_homem_ano', 'Inteiro', f"Horas-homem trabalhadas em 12 meses (colaboradores × {J['horasAno']} h)", FMT_INT, ''),
    ('grau_risco', 'Inteiro', 'Grau de risco da atividade (NR-4), usado no dimensionamento do SESMT', None, ''),
], linhas)

tabela('dSetor', 'Setores usados no inventário de riscos, brigada e pesquisa',
       [('setor', 'Texto (chave)', 'Nome do setor', None, '')], [[s] for s in J['setores']])

meses = []
for i in range(12):
    m = REF.month - 11 + i
    a = REF.year + (m - 1) // 12
    m = (m - 1) % 12 + 1
    meses.append(date(a, m, 1))
NOMES_MES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
tabela('dCalendario', 'Meses do período de 12 meses analisado', [
    ('mes', 'Data (chave)', 'Primeiro dia do mês', FMT_DATA, ''),
    ('ano', 'Inteiro', 'Ano', None, ''),
    ('numero_mes', 'Inteiro', 'Número do mês', None, ''),
    ('nome_mes', 'Texto', 'Mês abreviado (rótulo dos gráficos)', None, ''),
    ('semestre', 'Texto', 'Semestre do ano', None, ''),
], [[m, m.year, m.month, f'{NOMES_MES[m.month - 1]}/{str(m.year)[2:]}', f'{1 if m.month <= 6 else 2}º sem./{m.year}'] for m in meses])

# =====================================================================
# FATOS · ACIDENTES, RISCOS, BRIGADA E EQUIPAMENTOS
# =====================================================================
linhas = []
for u in UNIDADES:
    x = PU[u['id']]['d']
    for i, m in enumerate(meses):
        linhas.append([u['sigla'], m, x['mensalCAF'][i], x['mensalSAF'][i]])
tabela('fAcidentesMensal', 'Acidentes de trabalho por unidade e mês', [
    UN, ('mes', 'Data', 'Mês da ocorrência', FMT_DATA, 'dCalendario[mes]'),
    ('acidentes_caf', 'Inteiro', 'Acidentes com afastamento (CAF)', None, ''),
    ('acidentes_saf', 'Inteiro', 'Acidentes sem afastamento (SAF)', None, ''),
], linhas)

linhas = []
for u in UNIDADES:
    for setor in J['setores']:
        for i, cat in enumerate(J['categorias']):
            linhas.append([u['sigla'], setor, cat, PU[u['id']]['d']['riscos'][setor][i]])
tabela('fRiscos', 'Inventário de riscos do PGR (NR-1): perigos identificados por setor e categoria', [
    UN, ('setor', 'Texto', 'Setor', None, 'dSetor[setor]'),
    ('categoria_risco', 'Texto', 'Físico, Químico, Biológico, Ergonômico, Acidentes ou Psicossociais', None, ''),
    ('qtd_perigos', 'Inteiro', 'Quantidade de perigos identificados', None, ''),
], linhas)

SIT_BRIG = '=IF(H{r}<0,"Vencido",IF(H{r}<=30,"Vence em até 30 dias",IF(H{r}<=60,"Vence em até 60 dias","Válido")))'
linhas = []
for u in UNIDADES:
    for b in PU[u['id']]['d']['brigadistas']:
        linhas.append([b['id'], b['nome'], b['unidade'], b['setor'], b['turno'], d(b['treinamento']), d(b['validade']),
                       f'=G{{r}}-{REF_CELULA}', SIT_BRIG])
tabela('fBrigadistas', 'Brigadistas voluntários e validade do treinamento (NR-23)', [
    ('id_brigadista', 'Texto (chave)', 'Identificador do brigadista', None, ''),
    ('nome', 'Texto', 'Nome (fictício)', None, ''), UN,
    ('setor', 'Texto', 'Setor de lotação', None, 'dSetor[setor]'),
    ('turno', 'Texto', 'Turno de trabalho', None, ''),
    ('data_ultimo_treinamento', 'Data', 'Data do último treinamento de brigada', FMT_DATA, ''),
    ('data_validade', 'Data', 'Validade do treinamento (reciclagem anual)', FMT_DATA, ''),
    ('dias_para_vencimento', 'Fórmula', 'data_validade − data de referência (negativo = vencido)', FMT_INT, ''),
    ('situacao', 'Fórmula', 'Vencido, Vence em até 30 dias, Vence em até 60 dias ou Válido', None, ''),
], linhas, {'nome': 30, 'situacao': 22})

SIT_EQUIP = '=IF(H{r}<0,"Vencido",IF(H{r}<=30,"Até 30 dias",IF(H{r}<=90,"31 a 90 dias","Em dia")))'
linhas = []
for u in UNIDADES:
    for e in PU[u['id']]['d']['equipamentos']:
        linhas.append([e['codigo'], e['item'], e['categoria'], e['local'], e['unidade'], e['quantidade'], d(e['validade']),
                       f'=G{{r}}-{REF_CELULA}', SIT_EQUIP])
tabela('fEquipamentos', 'Extintores, hidrantes, EPIs e kits de emergência com validade', [
    ('codigo', 'Texto (chave)', 'Código patrimonial do item', None, ''),
    ('item', 'Texto', 'Descrição do item', None, ''),
    ('categoria', 'Texto', 'Extintor, Hidrante, EPI, Kit ou Emergência', None, ''),
    ('local', 'Texto', 'Local de instalação ou guarda', None, ''), UN,
    ('quantidade', 'Inteiro', 'Quantidade (lotes de EPI em unidades)', FMT_INT, ''),
    ('data_validade', 'Data', 'Validade ou próxima manutenção/recarga', FMT_DATA, ''),
    ('dias_para_vencimento', 'Fórmula', 'data_validade − data de referência', FMT_INT, ''),
    ('situacao', 'Fórmula', 'Vencido, Até 30 dias, 31 a 90 dias ou Em dia', None, ''),
], linhas)

SIT_DOC = '=IF(J{r}<0,"Vencido",IF(J{r}<=30,"Até 30 dias",IF(J{r}<=60,"Até 60 dias","Vigente")))'
linhas = []
for u in UNIDADES:
    for doc in PU[u['id']]['d']['documentos']:
        linhas.append([doc['codigo'], doc['nome'], doc['tipo'], doc['norma'], doc['unidade'], doc['responsavel'], doc['periodicidade'],
                       d(doc['ultimaRevisao']), d(doc['proximaRevisao']), f'=I{{r}}-{REF_CELULA}', SIT_DOC, doc['fluxo'], doc['progresso'] / 100])
tabela('fDocumentos', 'Documentos obrigatórios de SST, POPs e fluxos de trabalho (vigência e revisão)', [
    ('codigo', 'Texto', 'Código do documento (repete entre unidades)', None, ''),
    ('documento', 'Texto', 'Nome do documento', None, ''),
    ('tipo', 'Texto', 'Programa, Laudo, Plano, POP ou Fluxo', None, ''),
    ('norma', 'Texto', 'Norma de referência', None, ''), UN,
    ('responsavel', 'Texto', 'Área ou profissional responsável', None, ''),
    ('periodicidade', 'Texto', 'Periodicidade de revisão adotada', None, ''),
    ('data_ultima_revisao', 'Data', 'Data da última revisão', FMT_DATA, ''),
    ('data_proxima_revisao', 'Data', 'Data limite da próxima revisão', FMT_DATA, ''),
    ('dias_para_revisao', 'Fórmula', 'data_proxima_revisao − data de referência', FMT_INT, ''),
    ('vigencia', 'Fórmula', 'Vencido, Até 30 dias, Até 60 dias ou Vigente', None, ''),
    ('status_fluxo', 'Texto', 'Atualizado, Em revisão ou Desatualizado', None, ''),
    ('progresso_fluxo', 'Percentual', 'Avanço da atualização do fluxo ou POP', FMT_PCT, ''),
], linhas, {'documento': 44, 'periodicidade': 30})

# =====================================================================
# FATOS · PESQUISA DE PERCEPÇÃO
# =====================================================================
linhas = []
for u in UNIDADES:
    for x in PU[u['id']]['d']['pesquisa']['dimensoes']:
        linhas.append([u['sigla'], x['nome'], round(x['anterior'], 2), round(x['atual'], 2), '=D{r}-C{r}'])
tabela('fPesquisaDimensao', 'Pesquisa de clima de segurança: nota por dimensão (0 a 10)', [
    UN, ('dimensao', 'Texto', 'Dimensão avaliada', None, ''),
    ('nota_ciclo_anterior', 'Decimal', 'Nota média no ciclo anterior', FMT_DEC2, ''),
    ('nota_ciclo_atual', 'Decimal', 'Nota média no ciclo atual', FMT_DEC2, ''),
    ('variacao', 'Fórmula', 'nota_ciclo_atual − nota_ciclo_anterior', FMT_DEC2, ''),
], linhas)

linhas = [[u['sigla'], s, round(PU[u['id']]['d']['pesquisa']['porSetor'][s], 2)] for u in UNIDADES for s in J['setores']]
tabela('fPesquisaSetor', 'Pesquisa de clima de segurança: nota média por setor', [
    UN, ('setor', 'Texto', 'Setor', None, 'dSetor[setor]'), ('nota_media', 'Decimal', 'Nota média do setor (0 a 10)', FMT_DEC2, ''),
], linhas)

ano, sem, ciclos = REF.year, (1 if REF.month <= 6 else 2), []
for _ in range(4):
    ciclos.insert(0, f'{sem}º sem./{ano}')
    sem -= 1
    if sem == 0:
        sem, ano = 2, ano - 1
linhas = [[u['sigla'], i + 1, c, round(PU[u['id']]['d']['pesquisa']['tendencia'][i], 2)] for u in UNIDADES for i, c in enumerate(ciclos)]
tabela('fPesquisaCiclo', 'Evolução da nota geral da pesquisa nos últimos quatro ciclos', [
    UN, ('ordem_ciclo', 'Inteiro', '1 = mais antigo, 4 = ciclo atual', None, ''),
    ('ciclo', 'Texto', 'Semestre da aplicação', None, ''), ('nota_geral', 'Decimal', 'Nota geral do ciclo', FMT_DEC2, ''),
], linhas)

# =====================================================================
# FATO · INDICADORES DE GESTÃO (uma linha por unidade, 12 meses)
# =====================================================================
linhas = []
for u in UNIDADES:
    x, med = PU[u['id']]['d'], PU[u['id']]['med']
    ind = x['indicadores']
    linhas.append([u['sigla'], REF, x['diasPerdidos'], x['diasSemAcidente'], round(ind['conformidadeNR1'] / 100, 4), round(ind['trustIndex'] / 100, 4),
                   round(ind['horasTreinamento'], 1), ind['ncsAbertas'], round(ind['investigacoesNoPrazo'] / 100, 4), round(ind['absenteismo'] / 100, 4),
                   round(ind['turnover'] / 100, 4), x['brigadistasExigidos'], x['pesquisa']['respondentes'],
                   med['periodicosTotal'], med['periodicosVencidos'], med['periodicosAVencer'], med['desligamentos'], med['pppEmitidos']])
tabela('fIndicadoresGestao', 'Indicadores consolidados por unidade (janela de 12 meses até a data de referência)', [
    UN, ('data_referencia', 'Data', 'Data de fechamento dos indicadores', FMT_DATA, ''),
    ('dias_perdidos_12m', 'Inteiro', 'Dias perdidos e debitados por acidentes (taxa de gravidade)', FMT_INT, ''),
    ('dias_sem_acidente_caf', 'Inteiro', 'Dias desde o último acidente com afastamento', FMT_INT, ''),
    ('conformidade_nr1', 'Percentual', 'Itens conformes no checklist do GRO (NR-1)', FMT_PCT, ''),
    ('trust_index', 'Percentual', 'Respostas favoráveis na pesquisa de clima organizacional', FMT_PCT, ''),
    ('horas_treinamento_sst_colab', 'Decimal', 'Horas de treinamento em SST por colaborador no ano', FMT_DEC1, ''),
    ('ncs_abertas_auditoria', 'Inteiro', 'Não conformidades abertas na auditoria interna de SST', None, ''),
    ('investigacoes_no_prazo', 'Percentual', 'Investigações de incidentes concluídas no prazo', FMT_PCT, ''),
    ('absenteismo', 'Percentual', 'Taxa de absenteísmo', FMT_PCT, ''),
    ('turnover_voluntario', 'Percentual', 'Turnover voluntário em 12 meses', FMT_PCT, ''),
    ('brigadistas_exigidos', 'Inteiro', 'Brigadistas exigidos pelo dimensionamento da brigada', None, ''),
    ('respondentes_pesquisa', 'Inteiro', 'Respondentes da pesquisa de clima de segurança', FMT_INT, ''),
    ('periodicos_total', 'Inteiro', 'Colaboradores com ASO periódico previsto no PCMSO', FMT_INT, ''),
    ('periodicos_vencidos', 'Inteiro', 'ASOs periódicos vencidos', FMT_INT, ''),
    ('periodicos_a_vencer_30d', 'Inteiro', 'ASOs periódicos que vencem em até 30 dias', FMT_INT, ''),
    ('desligamentos_12m', 'Inteiro', 'Desligamentos em 12 meses', FMT_INT, ''),
    ('ppp_emitidos', 'Inteiro', 'Desligamentos com PPP eletrônico disponível', FMT_INT, ''),
], linhas)

# =====================================================================
# FATOS · ENGENHARIA E MEDICINA DO TRABALHO
# =====================================================================
linhas = [[u['sigla'], s['funcao'], s['exigido'], s['existente'], '=MAX(0,C{r}-D{r})'] for u in UNIDADES for s in PU[u['id']]['med']['sesmt']]
tabela('fSesmt', 'SESMT (NR-4): profissionais exigidos e existentes por unidade (dimensionamento fictício)', [
    UN, ('funcao', 'Texto', 'Profissional do SESMT', None, ''), ('exigido', 'Inteiro', 'Quantidade exigida', None, ''),
    ('existente', 'Inteiro', 'Quantidade em atuação', None, ''), ('deficit', 'Fórmula', 'MAX(0; exigido − existente)', None, ''),
], linhas)

linhas = [[u['sigla'], a['tipo'], a['noPrazo'], a['foraDoPrazo'], '=C{r}+D{r}'] for u in UNIDADES for a in PU[u['id']]['med']['asos']]
tabela('fAsos', 'ASOs realizados em 12 meses por tipo (PCMSO · NR-7)', [
    UN, ('tipo_aso', 'Texto', 'Admissional, Periódico, Retorno ao trabalho, Mudança de risco ou Demissional', None, ''),
    ('no_prazo', 'Inteiro', 'Realizados dentro do prazo', FMT_INT, ''), ('fora_do_prazo', 'Inteiro', 'Realizados fora do prazo', FMT_INT, ''),
    ('total', 'Fórmula', 'no_prazo + fora_do_prazo', FMT_INT, ''),
], linhas)

SIT_ASO = '=IF(G{r}<0,"Vencido","Vence em até 30 dias")'
linhas = [[p['nome'], p['unidade'], p['setor'], p['funcao'], d(p['ultimo']), d(p['vencimento']), f'=F{{r}}-{REF_CELULA}', SIT_ASO]
          for u in UNIDADES for p in PU[u['id']]['med']['pendentes']]
tabela('fPeriodicosPendentes', 'ASOs periódicos vencidos ou a vencer: casos prioritários (amostra)', [
    ('nome', 'Texto', 'Colaborador (fictício)', None, ''), UN, ('setor', 'Texto', 'Setor', None, 'dSetor[setor]'),
    ('funcao', 'Texto', 'Função', None, ''), ('data_ultimo_periodico', 'Data', 'Data do último ASO periódico', FMT_DATA, ''),
    ('data_vencimento', 'Data', 'Vencimento do periódico', FMT_DATA, ''),
    ('dias_para_vencimento', 'Fórmula', 'data_vencimento − data de referência', FMT_INT, ''),
    ('situacao', 'Fórmula', 'Vencido ou Vence em até 30 dias', None, ''),
], linhas, {'nome': 30, 'situacao': 22})

linhas = [[u['sigla'], e['exame'], e['programados'], e['realizados'], '=IF(C{r}=0,0,D{r}/C{r})'] for u in UNIDADES for e in PU[u['id']]['med']['exames']]
tabela('fExames', 'Exames complementares do PCMSO: programados × realizados', [
    UN, ('exame', 'Texto', 'Exame complementar', None, ''), ('programados', 'Inteiro', 'Exames programados no período', FMT_INT, ''),
    ('realizados', 'Inteiro', 'Exames realizados', FMT_INT, ''), ('pct_realizado', 'Fórmula', 'realizados ÷ programados', FMT_PCT, ''),
], linhas)

linhas = [[u['sigla'], g, PU[u['id']]['med']['afastamentos'][g]] for u in UNIDADES for g in J['gruposAfastamento']]
tabela('fAfastamentos', 'Dias de afastamento em 12 meses por grupo de causa (dados agregados)', [
    UN, ('grupo_causa', 'Texto', 'Grupo de causa do afastamento', None, ''), ('dias_afastamento', 'Inteiro', 'Dias de afastamento', FMT_INT, ''),
], linhas)

SIT_AG = '=IF(G{r}=0,"Controlado",IF(G{r}=1,"Em adequação","Sem controle eficaz"))'
linhas = [[u['sigla'], a['agente'], a['anexo'], ', '.join(a['setores']), a['expostos'], a['grau'], a['acimaDoLimite'], SIT_AG]
          for u in UNIDADES for a in PU[u['id']]['med']['agentes']]
tabela('fAgentesNocivos', 'Agentes nocivos e insalubridade identificados no LTCAT (NR-15)', [
    UN, ('agente', 'Texto', 'Agente nocivo', None, ''), ('referencia_nr15', 'Texto', 'Anexo da NR-15', None, ''),
    ('setores_expostos', 'Texto', 'Setores com exposição', None, ''), ('trabalhadores_expostos', 'Inteiro', 'Trabalhadores expostos', FMT_INT, ''),
    ('grau_insalubridade', 'Texto', 'Grau de insalubridade e adicional', None, ''),
    ('medicoes_acima_limite', 'Inteiro', 'Medições acima do limite de tolerância no período', None, ''),
    ('situacao_controle', 'Fórmula', 'Controlado (0), Em adequação (1) ou Sem controle eficaz (2 ou mais)', None, ''),
], linhas)

linhas = [[u['sigla'], e['codigo'], e['descricao'], e['prazo'], e['enviados'], e['pendentes'], e['rejeitados']] for u in UNIDADES for e in PU[u['id']]['med']['eventos']]
tabela('fEventosEsocial', 'Eventos de SST enviados ao eSocial (base do PPP eletrônico)', [
    UN, ('evento', 'Texto', 'S-2210 (CAT), S-2220 (ASO) ou S-2240 (agentes nocivos)', None, ''), ('descricao', 'Texto', 'Descrição do evento', None, ''),
    ('prazo_envio', 'Texto', 'Prazo legal de envio', None, ''), ('enviados', 'Inteiro', 'Eventos enviados e aceitos', FMT_INT, ''),
    ('pendentes', 'Inteiro', 'Eventos pendentes de envio', FMT_INT, ''), ('rejeitados', 'Inteiro', 'Eventos rejeitados', FMT_INT, ''),
], linhas, {'prazo_envio': 46})

# =====================================================================
# METAS DOS INDICADORES (selos e certificações)
# =====================================================================
metas = J['metricasConsolidado']
PCT = {'%'}
linhas = []
for m in metas:
    eh_pct = m['unidade'] == '%'
    meta = 3 if m['id'] == 'ncs' else (m['meta'] / 100 if eh_pct else m['meta'])
    unidade_med = {'%': 'Percentual', 'h': 'Horas por colaborador', '': 'Número'}[m['unidade']]
    if m['id'] == 'psi':
        unidade_med = 'Nota de 0 a 10'
    if m['id'] == 'ncs':
        unidade_med = 'Quantidade por unidade (no consolidado: 3 × nº de unidades)'
    linhas.append([m['id'], m['nome'], meta, 'Maior ou igual' if m['sentido'] == 'maior' else 'Menor ou igual', unidade_med,
                   ', '.join(m['selos']), ', '.join(m.get('bloqueante') or []), m['sugestao']])
ws = tabela('dMetas', 'Metas internas de prontidão dos indicadores e selos afetados (simulação)', [
    ('id_indicador', 'Texto (chave)', 'Identificador do indicador', None, ''), ('indicador', 'Texto', 'Nome do indicador', None, ''),
    ('meta', 'Número', 'Valor da meta (percentuais em fração: 0,9 = 90%)', None, ''), ('sentido', 'Texto', 'Se a meta é mínima ou máxima', None, ''),
    ('unidade_medida', 'Texto', 'Unidade da meta', None, ''), ('selos_afetados', 'Texto', 'GPTW, ISO 45001 e/ou ESG', None, ''),
    ('bloqueante_para', 'Texto', 'Selo que fica inelegível se a meta não for atingida', None, ''),
    ('sugestao_melhoria', 'Texto', 'Plano de mitigação sugerido quando fora da meta', None, ''),
], linhas, {'indicador': 46, 'sugestao_melhoria': 80, 'unidade_medida': 30})
for r in range(2, len(linhas) + 2):
    cel = ws.cell(row=r, column=3)
    cel.font = FONTE_ENTRADA
    cel.number_format = FMT_PCT if linhas[r - 2][4] == 'Percentual' else '0.0#'
    ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical='top')

# =====================================================================
# CONFERÊNCIA: KPIs recalculados a partir das tabelas × valores do painel
# =====================================================================
k = J['kpisConsolidado']
med_tot = {c: sum(PU[u['id']]['med'][c] for u in UNIDADES) for c in ('periodicosTotal', 'periodicosVencidos', 'desligamentos', 'pppEmitidos')}
exig = sum(s['exigido'] for u in UNIDADES for s in PU[u['id']]['med']['sesmt'])
cob = sum(min(s['existente'], s['exigido']) for u in UNIDADES for s in PU[u['id']]['med']['sesmt'])
n_dim = len(J['dimensoes'])
ws = wb.create_sheet('Conferencia')
ws['A1'] = 'Conferência: indicadores recalculados a partir das tabelas × valores exibidos no painel (consolidado, todas as unidades)'
ws['A1'].font = FONTE_TITULO
ws['A2'] = 'Coluna C: fórmulas sobre as tabelas desta base. Coluna D (azul): valor exibido no painel HTML, copiado do relatório. Coluna E: diferença.'
ws['A2'].font = Font(name='Arial', size=9, italic=True, color='605E5C')
cab = ['Indicador', 'Regra de cálculo', 'Calculado na base', 'Valor no painel', 'Diferença']
for c, t in enumerate(cab, 1):
    ws.cell(row=4, column=c, value=t).font = FONTE_NEGRITO
    ws.cell(row=4, column=c).fill = PatternFill('solid', fgColor='DCE6F1')
CONF = [
    ('Colaboradores', 'Soma de dUnidade[colaboradores]', '=SUM(dUnidade!F2:F5)', sum(u['colaboradores'] for u in UNIDADES), FMT_INT),
    ('Taxa de frequência (CAF)', 'Acidentes CAF ÷ horas-homem × 1.000.000', '=SUM(fAcidentesMensal!C:C)/SUM(dUnidade!G2:G5)*1000000', k['tf'], FMT_DEC2),
    ('Taxa de gravidade', 'Dias perdidos ÷ horas-homem × 1.000.000', '=SUM(fIndicadoresGestao!C2:C5)/SUM(dUnidade!G2:G5)*1000000', k['tg'], FMT_DEC1),
    ('Conformidade NR-1', 'Média ponderada pelo efetivo', '=SUMPRODUCT(fIndicadoresGestao!E2:E5,SUMIF(dUnidade!A2:A5,fIndicadoresGestao!A2:A5,dUnidade!F2:F5))/SUM(dUnidade!F2:F5)', k['conformidadeNR1'] / 100, FMT_PCT),
    ('Brigadistas com treinamento válido', 'Não vencidos ÷ exigidos (limitado a 100%)', '=MIN(1,(COUNTA(fBrigadistas!A:A)-1-COUNTIF(fBrigadistas!I:I,"Vencido"))/SUM(fIndicadoresGestao!L2:L5))', k['indiceBrigada'] / 100, FMT_PCT),
    ('Equipamentos vencidos', 'Contagem de situação = Vencido', '=COUNTIF(fEquipamentos!I:I,"Vencido")', k['equipVencidos'], FMT_INT),
    ('Documentos fora da vigência', 'Contagem de vigência = Vencido', '=COUNTIF(fDocumentos!K:K,"Vencido")', k['docsVencidos'], FMT_INT),
    ('Nota de percepção de segurança', 'Média das dimensões ponderada pelos respondentes', f'=SUMPRODUCT(fPesquisaDimensao!D2:D{4 * n_dim + 1},SUMIF(fIndicadoresGestao!A2:A5,fPesquisaDimensao!A2:A{4 * n_dim + 1},fIndicadoresGestao!M2:M5))/(SUM(fIndicadoresGestao!M2:M5)*{n_dim})', k['notaPercepcao'], FMT_DEC2),
    ('Cobertura do SESMT', '(Exigido − déficit) ÷ exigido', '=(SUM(fSesmt!C:C)-SUM(fSesmt!E:E))/SUM(fSesmt!C:C)', cob / exig, FMT_PCT),
    ('ASOs periódicos em dia', '1 − vencidos ÷ total previsto', '=1-SUM(fIndicadoresGestao!O2:O5)/SUM(fIndicadoresGestao!N2:N5)', 1 - med_tot['periodicosVencidos'] / med_tot['periodicosTotal'], FMT_PCT),
    ('PPP disponível nos desligamentos', 'PPP emitidos ÷ desligamentos', '=SUM(fIndicadoresGestao!R2:R5)/SUM(fIndicadoresGestao!Q2:Q5)', med_tot['pppEmitidos'] / med_tot['desligamentos'], FMT_PCT),
]
for i, (nome, regra, formula, valor, fmt) in enumerate(CONF, 5):
    ws.cell(row=i, column=1, value=nome).font = FONTE
    ws.cell(row=i, column=2, value=regra).font = FONTE
    c = ws.cell(row=i, column=3, value=formula); c.font = FONTE; c.number_format = fmt
    c = ws.cell(row=i, column=4, value=valor); c.font = FONTE_ENTRADA; c.number_format = fmt
    c = ws.cell(row=i, column=5, value=f'=C{i}-D{i}'); c.font = FONTE; c.number_format = '0.0000'
fim = 4 + len(CONF)
ws.cell(row=fim + 2, column=1, value='Resultado da conferência').font = FONTE_NEGRITO
ws.cell(row=fim + 2, column=3, value=f'=IF(SUMPRODUCT(ABS(E5:E{fim}))<0.01,"Base confere com o painel","Verificar diferenças")').font = FONTE_NEGRITO
ws.cell(row=fim + 3, column=1, value='Observação: notas e percentuais das tabelas estão arredondados (2 e 4 casas); diferenças abaixo de 0,01 vêm só desse arredondamento.').font = Font(name='Arial', size=9, italic=True, color='605E5C')
for col, w in zip('ABCDE', (36, 48, 20, 18, 14)):
    ws.column_dimensions[col].width = w

# =====================================================================
# DICIONÁRIO DE DADOS
# =====================================================================
ws = wb.create_sheet('Dicionario')
for c, t in enumerate(['tabela', 'coluna', 'tipo', 'descricao', 'relacionamento'], 1):
    ws.cell(row=1, column=c, value=t).font = FONTE_NEGRITO
for i, linha in enumerate(DICIONARIO, 2):
    for c, v in enumerate(linha, 1):
        ws.cell(row=i, column=c, value=v).font = FONTE
t = Table(displayName='Dicionario', ref=f'A1:E{len(DICIONARIO) + 1}')
t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
ws.add_table(t)
ws.freeze_panes = 'A2'
for col, w in zip('ABCDE', (22, 30, 16, 80, 26)):
    ws.column_dimensions[col].width = w

# =====================================================================
# LEIA-ME (primeira aba)
# =====================================================================
ws = wb['Sheet']
ws.title = 'LEIA-ME'
ws['A1'] = 'Base de dados · Painel de SST e Conformidade'
ws['A1'].font = FONTE_TITULO
ws['A2'] = 'Todos os nomes, unidades e números são fictícios e foram gerados para demonstração.'
ws['A2'].font = Font(name='Arial', size=10, italic=True, color='A4262C')
ws['A4'] = 'Data de referência'
ws['A4'].font = FONTE_NEGRITO
ws['B4'] = REF
ws['B4'].number_format = FMT_DATA
ws['B4'].font = FONTE_ENTRADA
ws['B4'].fill = PatternFill('solid', fgColor='FFFF00')
ws['C4'] = 'Célula editável. Prazos e situações (vencido, a vencer, vigente) são calculados a partir desta data. Para acompanhar o dia corrente, troque por =HOJE().'
ws['C4'].font = Font(name='Arial', size=9, color='605E5C')
ws['A6'] = 'Convenções'
ws['A6'].font = FONTE_NEGRITO
ws['A7'] = 'Prefixo d = dimensão (cadastro); prefixo f = fato (medições). Chave principal entre tabelas: id_unidade (sigla da unidade).'
ws['A8'] = 'Percentuais estão em fração (0,9 = 90%). Colunas do tipo "Fórmula" são recalculadas pelo Excel e não devem ser digitadas.'
ws['A9'] = 'Valores em azul foram digitados; fundo amarelo indica célula de parâmetro.'
for r in (7, 8, 9):
    ws[f'A{r}'].font = FONTE
ws['A11'] = 'Tabelas'
ws['A11'].font = FONTE_NEGRITO
for c, t in enumerate(['Aba / tabela', 'Conteúdo', 'Linhas'], 1):
    ws.cell(row=12, column=c, value=t).font = FONTE_NEGRITO
    ws.cell(row=12, column=c).fill = PatternFill('solid', fgColor='DCE6F1')
for i, (nome, desc, n) in enumerate(TABELAS, 13):
    ws.cell(row=i, column=1, value=nome).font = FONTE
    ws.cell(row=i, column=2, value=desc).font = FONTE
    ws.cell(row=i, column=3, value=n).font = FONTE
proxima = 13 + len(TABELAS)
ws.cell(row=proxima, column=1, value='Conferencia').font = FONTE
ws.cell(row=proxima, column=2, value='KPIs recalculados com fórmulas a partir das tabelas e comparados com o painel').font = FONTE
ws.cell(row=proxima + 1, column=1, value='Dicionario').font = FONTE
ws.cell(row=proxima + 1, column=2, value='Descrição de cada coluna, tipo e relacionamento').font = FONTE
ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 90
ws.column_dimensions['C'].width = 10

# Recalcula todas as fórmulas ao abrir no Excel
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(SAIDA)
print('ok', len(TABELAS), 'tabelas,', len(DICIONARIO), 'colunas no dicionário')
